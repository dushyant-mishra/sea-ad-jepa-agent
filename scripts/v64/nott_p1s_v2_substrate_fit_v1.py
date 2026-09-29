#!/usr/bin/env python3
"""P1S V2 — Nott same-study microglial substrate fit, peak-opportunity adjusted.

Executes V64_NOTT_P1S_OPERATIONAL_SUCCESSOR_CONTRACT_V2 exactly. Nothing here is
chosen by me: the support rule, the null construction, the estimator, the
bootstrap replicate count and seed, and the three pass requirements are all
frozen in that contract.

WHY V2 SUPERSEDED V1: a raw cross-cell-type support-rate contrast does not
separate contact-specific alignment from cell-type differences in ATAC peak
opportunity. V2's primary comparison is a difference-in-differences --
observed-minus-null WITHIN each track, then microglial excess against comparator
excess -- so a track with more peaks overall gains nothing automatically.

THE NULL IS DETERMINISTIC, NOT A SHUFFLE. The promoter is held fixed and the
distal interval is mirrored to signed displacement -d with identical width. That
sidesteps the replicate-count question entirely: there is exactly one null
opportunity per constructible observed edge, so no Monte Carlo precision
authority is needed and none is invented.

ORIENTATION IS RESOLVED IN SOURCE hg19 COORDINATES, before liftover, so the
coordinate transform cannot determine promoter identity. The orientation audit is
recomputed here independently rather than taken on report.

TRAINING=OFF. TD60=BLOCKED. No external gene annotation, no target or AD identity,
no project RNA. P1S establishes INTERNAL substrate compatibility only.
"""
from __future__ import annotations

import argparse
import gzip
import hashlib
import json
import os
import sys
from bisect import bisect_right
from collections import Counter, defaultdict

import numpy as np
import openpyxl

S5 = "C:/Users/dushy/Downloads/NIHMS1066836-supplement-Table_S5.xlsx"
S5_SHA = "81c99689533d9da372cecdd469e7ff02cc985720105b83b3bd66c3ac8c93972e"
C3 = "results/v64/V64_NOTT_C3_RETAINED_CONTACTS_HG38.tsv.gz"
DISPO = "results/v64/c3_intermediates/V64_C3_PER_ANCHOR_DISPOSITION.tsv.gz"
ATAC = {"microglia": "C:/Users/dushy/jepa_c3/ATAC_PU1.bed.gz",
        "neuron": "C:/Users/dushy/jepa_c3/ATAC_NeuN.bed.gz",
        "oligo": "C:/Users/dushy/jepa_c3/ATAC_Olig2.bed.gz"}
CHAIN = "C:/Users/dushy/jepa_c3/hg19ToHg38.over.chain.gz"
N_BOOT = 4000            # frozen
BOOT_SEED = 20260929     # frozen


def sha(p):
    return hashlib.sha256(open(p, "rb").read()).hexdigest()


class Intervals:
    """Sorted per-chromosome intervals with >=1 bp overlap query."""

    def __init__(self):
        self.d = defaultdict(list)
        self._ready = False

    def add(self, c, s, e):
        self.d[c].append((s, e))

    def build(self):
        self.starts, self.ends = {}, {}
        for c, v in self.d.items():
            v.sort()
            self.starts[c] = [x[0] for x in v]
            # running max end, so a single bisect answers the query
            m, run = [], -1
            for _, e in v:
                run = max(run, e)
                m.append(run)
            self.ends[c] = m
        self._ready = True

    def overlaps(self, c, s, e) -> bool:
        if not self._ready or c not in self.starts:
            return False
        st = self.starts[c]
        i = bisect_right(st, e - 1)          # intervals starting before e
        if i == 0:
            return False
        return self.ends[c][i - 1] > s       # any of them ending after s

    def midpoints(self, c):
        return [(a + b) // 2 for a, b in sorted(self.d.get(c, []))]


def hg38_chrom_sizes(chain_path):
    """qSize from the hg19->hg38 chain headers: authoritative hg38 lengths from an
    already-authenticated file rather than a new download."""
    sizes = {}
    with gzip.open(chain_path, "rt") as fh:
        for line in fh:
            if line.startswith("chain"):
                f = line.split()
                sizes[f[7]] = int(f[8])
    return sizes


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--structure-only", action="store_true",
                    help="Stop BEFORE any ATAC support is computed. Orientation and the "
                         "population funnel are source-structure facts and open no P1S "
                         "outcome; the support rates ARE the outcome.")
    a = ap.parse_args()
    if os.path.exists(a.out_dir) and os.listdir(a.out_dir):
        raise SystemExit("STOP_OUTPUT_EXISTS")
    os.makedirs(a.out_dir, exist_ok=True)
    if sha(S5) != S5_SHA:
        raise SystemExit("STOP_S5_DIGEST_MISMATCH")

    # ---------- 1. active promoter annotation, source hg19
    wb = openpyxl.load_workbook(S5, read_only=True, data_only=True)
    ws = wb["H3K4me3_around_TSS_annotated_pe"]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows)
              if r and any(str(x).strip() == "PeakID" for x in r if x))
    hdr = [str(x).strip() if x is not None else "" for x in rows[hi]]
    col = {k: i for i, k in enumerate(hdr)}
    need = ["Chr", "Start", "End", "PU1_active_promoter", "Nearest Ensembl"]
    miss = [k for k in need if k not in col]
    if miss:
        raise SystemExit(f"STOP_ANNOTATION_SCHEMA {miss} in {hdr}")

    prom = Intervals()
    prom_rows = 0
    prom_by_chrom = defaultdict(list)          # (s,e,ensembl)
    for r in rows[hi + 1:]:
        if not r or r[col["Chr"]] is None:
            continue
        flag = r[col["PU1_active_promoter"]]
        if not (flag is True or str(flag).strip().upper() in ("TRUE", "1", "YES")):
            continue
        c, s, e = str(r[col["Chr"]]), int(r[col["Start"]]), int(r[col["End"]])
        ens = r[col["Nearest Ensembl"]]
        ens = str(ens).strip() if ens is not None else ""
        prom.add(c, s, e)
        prom_by_chrom[c].append((s, e, ens))
        prom_rows += 1
    prom.build()
    for c in prom_by_chrom:
        prom_by_chrom[c].sort()

    # ---------- 2. interactions, source hg19, orientation recomputed independently
    ws = wb["Microglia interactome"]
    irows = list(ws.iter_rows(values_only=True))
    ih = next(i for i, r in enumerate(irows) if r and r[0] == "chr1")
    inter = []
    for r in irows[ih + 1:]:
        if not r or r[0] is None:
            continue
        inter.append((str(r[0]), int(r[1]), int(r[2]),
                      str(r[3]), int(r[4]), int(r[5])))
    N0 = len(inter)

    def prom_hits(c, s, e):
        out = []
        for ps, pe, ens in prom_by_chrom.get(c, []):
            if ps >= e:
                break
            if pe > s:
                out.append(ens)
        return out

    orient = {}
    cls = Counter()
    for i, (c1, s1, e1, c2, s2, e2) in enumerate(inter):
        h1, h2 = prom_hits(c1, s1, e1), prom_hits(c2, s2, e2)
        p1, p2 = len(h1) > 0, len(h2) > 0
        if p1 and p2:
            cls["PROMOTER_PROMOTER"] += 1
            orient[i] = ("PROMOTER_PROMOTER", None)
        elif not p1 and not p2:
            cls["NO_ACTIVE_PROMOTER_MATCH"] += 1
            orient[i] = ("NO_ACTIVE_PROMOTER_MATCH", None)
        else:
            side = "A" if p1 else "B"
            genes = {g for g in (h1 if p1 else h2) if g}
            if len(genes) == 1:
                k = "PROMOTER_DISTAL_UNIQUE_ENSEMBL"
            elif len(genes) > 1:
                k = "GENE_AMBIGUOUS_PROMOTER_DISTAL"
            else:
                k = "GENE_UNANNOTATED_PROMOTER_DISTAL"
            cls[k] += 1
            cls["PROMOTER_DISTAL_EXACTLY_ONE_PROMOTER_ANCHOR_total"] += 1
            orient[i] = (k, side)

    # ---------- 3. C3-admissible pairs, with hg38 coordinates
    c3 = {}
    keep_idx = []
    with gzip.open(DISPO, "rt") as fh:
        h = fh.readline().rstrip("\n").split("\t")
        ci = {k: j for j, k in enumerate(h)}
        cur = {}
        for line in fh:
            f = line.rstrip("\n").split("\t")
            if f[ci["pair_C3_retained"]] != "True":
                continue
            i = int(f[ci["pair_index"]])
            cur.setdefault(i, {})[f[ci["side"]]] = (
                f[ci["hg38_chrom"]], int(f[ci["hg38_start"]]), int(f[ci["hg38_end"]]))
    for i, d in cur.items():
        if "A" in d and "B" in d:
            c3[i] = d
            keep_idx.append(i)

    # ---------- 4. P1S population: C3-admissible AND source-oriented promoter-distal
    P1S_OK = {"PROMOTER_DISTAL_UNIQUE_ENSEMBL", "GENE_AMBIGUOUS_PROMOTER_DISTAL",
              "GENE_UNANNOTATED_PROMOTER_DISTAL"}
    pop = [i for i in keep_idx if orient[i][0] in P1S_OK]

    if a.structure_only:
        out = {
            "schema": "V64_NOTT_P1S_STRUCTURE_ONLY_V1",
            "date": "2026-09-29",
            "status": "SOURCE_STRUCTURE_AND_POPULATION_FUNNEL_ONLY__NO_P1S_OUTCOME_OPENED",
            "why_stopped_here": (
                "V64_NOTT_P1S_OPERATIONAL_SUCCESSOR_CONTRACT_V2 is labelled "
                "'current_parallel_P1S_operational_authority_candidate' by "
                "V64_SPILLOVER_FIREWALL_MANIFEST_V1 and is absent from the canonical "
                "authority branch at 2d6fa8a1. Computing ATAC support would open the "
                "P1S outcome against a CANDIDATE contract and would irreversibly "
                "consume its prospectivity: once the support rates are seen, that "
                "contract can no longer be revised prospectively. Orientation and the "
                "population funnel are source-structure facts and open nothing."),
            "governance": {"training": "OFF", "td60": "BLOCKED",
                           "p1s_outcome_opened": False,
                           "atac_support_computed": False,
                           "external_gene_annotation_used": False,
                           "target_or_AD_identity_used": False},
            "authority_binding": {
                "canonical_branch": "chatgpt/v64-e2-single-source-successor-20260929",
                "canonical_commit": "2d6fa8a1c4201e836e089b798d36617c07d92a7e",
                "this_branch_descends_from_canonical": True,
                "p1s_v2_contract_status": "CANDIDATE, on chatgpt/v64-parallel-execution-readiness-20260929, NOT on canonical",
                "table_s5_sha256": sha(S5)},
            "orientation_recomputed_independently": {
                "original_interactions": N0,
                "PU1_active_promoter_annotation_rows": prom_rows,
                "counts": dict(cls),
                "matches_parallel_lane_audit": {
                    "PROMOTER_DISTAL_exactly_one": cls["PROMOTER_DISTAL_EXACTLY_ONE_PROMOTER_ANCHOR_total"] == 64210,
                    "PROMOTER_PROMOTER": cls["PROMOTER_PROMOTER"] == 7614,
                    "NO_ACTIVE_PROMOTER_MATCH": cls["NO_ACTIVE_PROMOTER_MATCH"] == 32978,
                    "PROMOTER_DISTAL_UNIQUE_ENSEMBL": cls["PROMOTER_DISTAL_UNIQUE_ENSEMBL"] == 61624}},
            "population_funnel": {
                "original_interactions": N0,
                "C3_retained": len(keep_idx),
                "C3_retained_and_source_oriented_promoter_distal": len(pop),
                "fraction_of_original": round(len(pop) / N0, 6)},
        }
        with open(os.path.join(a.out_dir, "V64_NOTT_P1S_STRUCTURE_ONLY_V1.json"), "w") as fh:
            json.dump(out, fh, indent=2)
        print("orientation (recomputed independently):", dict(cls))
        print("matches parallel audit:", out["orientation_recomputed_independently"]["matches_parallel_lane_audit"])
        print("funnel:", out["population_funnel"])
        print("")
        print("STOPPED BEFORE ATAC SUPPORT. No P1S outcome opened.")
        return 0

    # ---------- 5. ATAC tracks (hg38 columns 1-3)
    tracks = {}
    for name, path in ATAC.items():
        iv = Intervals()
        with gzip.open(path, "rt") as fh:
            for line in fh:
                if line.startswith("#") or not line.strip():
                    continue
                f = line.rstrip("\n").split("\t")
                iv.add(f[0], int(f[1]), int(f[2]))
        iv.build()
        tracks[name] = iv
    sizes = hg38_chrom_sizes(CHAIN)

    # ---------- 6. per-pair observed/null support
    rec = []
    null_bad = 0
    for i in pop:
        side = orient[i][1]
        pc, ps, pe = c3[i][side]
        dc, ds, de = c3[i]["B" if side == "A" else "A"]
        if pc != dc:
            null_bad += 1
            continue
        pm, dm, w = (ps + pe) // 2, (ds + de) // 2, de - ds
        d = dm - pm
        nm = pm - d
        ns, ne = nm - w // 2, nm - w // 2 + w
        if ns < 0 or ne > sizes.get(dc, 0):
            null_bad += 1
            continue
        r = {"i": i, "prom": (pc, ps, pe)}
        for t, iv in tracks.items():
            r["O_" + t] = int(iv.overlaps(dc, ds, de))
            r["N_" + t] = int(iv.overlaps(dc, ns, ne))
        rec.append(r)

    n = len(rec)
    if n == 0:
        raise SystemExit("STOP_EMPTY_P1S_POPULATION")
    E = {t: np.array([r["O_" + t] - r["N_" + t] for r in rec], float) for t in tracks}
    O = {t: np.array([r["O_" + t] for r in rec], float) for t in tracks}
    Nn = {t: np.array([r["N_" + t] for r in rec], float) for t in tracks}
    D_MN = E["microglia"] - E["neuron"]
    D_MO = E["microglia"] - E["oligo"]

    # ---------- 7. promoter-cluster bootstrap, frozen replicates and seed
    prom_key = [r["prom"] for r in rec]
    uniq = sorted(set(prom_key))
    idx_of = {p: j for j, p in enumerate(uniq)}
    groups = defaultdict(list)
    for j, p in enumerate(prom_key):
        groups[idx_of[p]].append(j)
    glist = [np.array(groups[j]) for j in range(len(uniq))]
    rng = np.random.default_rng(BOOT_SEED)

    def boot_ci(v):
        means = np.empty(N_BOOT)
        for b in range(N_BOOT):
            pick = rng.integers(0, len(glist), len(glist))
            sel = np.concatenate([glist[k] for k in pick])
            means[b] = v[sel].mean()
        lo, hi = np.quantile(means, [0.025, 0.975])
        return float(v.mean()), float(lo), float(hi)

    res = {}
    for label, v in (("E_microglia", E["microglia"]), ("E_neuron", E["neuron"]),
                     ("E_oligo", E["oligo"]), ("D_MN", D_MN), ("D_MO", D_MO)):
        m, lo, hi = boot_ci(v)
        res[label] = {"mean": m, "ci95_low": lo, "ci95_high": hi}

    R1 = res["E_microglia"]["ci95_low"] > 0
    R2 = res["D_MN"]["ci95_low"] > 0
    R3 = res["D_MO"]["ci95_low"] > 0

    # ---------- 8. four-way counts and density sensitivity
    def fourway(x, y):
        return {"both": int(((x == 1) & (y == 1)).sum()),
                "microglia_only": int(((x == 1) & (y == 0)).sum()),
                "comparator_only": int(((x == 0) & (y == 1)).sum()),
                "neither": int(((x == 0) & (y == 0)).sum())}

    dens = {}
    mids = {t: {c: np.array(iv.midpoints(c)) for c in iv.d} for t, iv in tracks.items()}
    for t in tracks:
        cnt = []
        for r in rec:
            dc = c3[r["i"]]["B" if orient[r["i"]][1] == "A" else "A"][0]
            m = (c3[r["i"]]["B" if orient[r["i"]][1] == "A" else "A"][1]
                 + c3[r["i"]]["B" if orient[r["i"]][1] == "A" else "A"][2]) // 2
            arr = mids[t].get(dc, np.array([]))
            cnt.append(int(np.searchsorted(arr, m + 50000) - np.searchsorted(arr, m - 50000)))
        dens[t] = np.array(cnt, float)
    q = np.quantile(dens["microglia"], [0.25, 0.5, 0.75])
    strat = np.digitize(dens["microglia"], q)
    by_density = {}
    for s in range(4):
        m = strat == s
        if m.sum() < 2:
            continue
        by_density[f"Q{s+1}"] = {"n": int(m.sum()),
                                 "E_microglia": float(E["microglia"][m].mean()),
                                 "D_MN": float(D_MN[m].mean()),
                                 "D_MO": float(D_MO[m].mean())}

    out = {
        "schema": "V64_NOTT_P1S_V2_SUBSTRATE_FIT_RESULT_V1",
        "date": "2026-09-29",
        "contract": "V64_NOTT_P1S_OPERATIONAL_SUCCESSOR_CONTRACT_V2",
        "governance": {"training": "OFF", "td60": "BLOCKED",
                       "external_gene_annotation_used": False,
                       "target_or_AD_identity_used": False,
                       "continuous_adjustment_estimator_imported": False},
        "inputs": {"table_s5_sha256": sha(S5),
                   "c3_retained_sha256": sha(C3),
                   "atac_sha256": {k: sha(v) for k, v in ATAC.items()}},
        "orientation_recomputed_independently": {
            "original_interactions": N0,
            "PU1_active_promoter_annotation_rows": prom_rows,
            "counts": dict(cls),
            "matches_parallel_lane_audit": {
                "PROMOTER_DISTAL_exactly_one": cls["PROMOTER_DISTAL_EXACTLY_ONE_PROMOTER_ANCHOR_total"] == 64210,
                "PROMOTER_PROMOTER": cls["PROMOTER_PROMOTER"] == 7614,
                "NO_ACTIVE_PROMOTER_MATCH": cls["NO_ACTIVE_PROMOTER_MATCH"] == 32978,
                "PROMOTER_DISTAL_UNIQUE_ENSEMBL": cls["PROMOTER_DISTAL_UNIQUE_ENSEMBL"] == 61624}},
        "population_funnel": {
            "original": N0,
            "C3_retained": len(keep_idx),
            "C3_retained_and_source_oriented_promoter_distal": len(pop),
            "NULL_NOT_CONSTRUCTIBLE": null_bad,
            "NULL_NOT_CONSTRUCTIBLE_fraction": round(null_bad / max(1, len(pop)), 6),
            "analysed_pairs": n,
            "distinct_promoter_clusters": len(uniq)},
        "support_rates": {t: {"observed": float(O[t].mean()),
                              "null": float(Nn[t].mean()),
                              "excess_E": float(E[t].mean())} for t in tracks},
        "primary": res,
        "requirements": {"R1_microglia_above_own_null": bool(R1),
                         "R2_microglia_excess_above_neuron": bool(R2),
                         "R3_microglia_excess_above_oligo": bool(R3)},
        "four_way_observed": {
            "microglia_vs_neuron": fourway(O["microglia"], O["neuron"]),
            "microglia_vs_oligo": fourway(O["microglia"], O["oligo"])},
        "four_way_null": {
            "microglia_vs_neuron": fourway(Nn["microglia"], Nn["neuron"]),
            "microglia_vs_oligo": fourway(Nn["microglia"], Nn["oligo"])},
        "local_density_sensitivity_descriptive_only": by_density,
        "bootstrap": {"replicates": N_BOOT, "seed": BOOT_SEED,
                      "cluster": "promoter anchor identity",
                      "interval": "two-sided 95% percentile"},
        "claim_scope": ("A P1S pass establishes INTERNAL microglial substrate "
                        "compatibility of the Nott-centered contact object relative to "
                        "same-study neuronal/oligodendrocyte ATAC and the frozen "
                        "geometry null. It is NOT independent external corroboration, "
                        "donor-independent replication, multi-source contact "
                        "replication, or causal enhancer-gene assignment."),
    }
    out["P1S_ADJUDICATION"] = "PASS" if (R1 and R2 and R3) else "FAIL"
    if not out["P1S_ADJUDICATION"] == "PASS":
        out["hard_stop"] = ("NOT_MICROGLIA_FIT_FOR_TARGET_CONSTRUCTION under P1S. A "
                            "primary fail may not be rescued by a favourable density "
                            "stratum.")
    with open(os.path.join(a.out_dir, "V64_NOTT_P1S_V2_RESULT_V1.json"), "w") as fh:
        json.dump(out, fh, indent=2)

    o = out
    print("orientation (recomputed):", dict(cls))
    print("matches parallel audit:", o["orientation_recomputed_independently"]["matches_parallel_lane_audit"])
    print("funnel:", o["population_funnel"])
    print("support rates:")
    for t, v in o["support_rates"].items():
        print(f"   {t:10s} observed {v['observed']:.4f}  null {v['null']:.4f}  E {v['excess_E']:+.4f}")
    print("primary:")
    for k, v in res.items():
        print(f"   {k:12s} mean {v['mean']:+.5f}  CI95 [{v['ci95_low']:+.5f}, {v['ci95_high']:+.5f}]")
    print(f"R1 {R1}  R2 {R2}  R3 {R3}")
    print(f"P1S_ADJUDICATION: {out['P1S_ADJUDICATION']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
