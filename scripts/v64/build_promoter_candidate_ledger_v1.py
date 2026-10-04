#!/usr/bin/env python3
"""Build an annotation-first promoter/TSS candidate ledger.

GENCODE defines the candidate denominator. SCREEN, Dong and FANTOM are evidence
layers only and may not remove candidates.
"""
from __future__ import annotations

import argparse
import bisect
import csv
import gzip
import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path

ATTR_RE = re.compile(r'(\S+) "([^"]+)"')


def strip_version(x: str) -> str:
    return x.split(".", 1)[0] if x else x


def load_gencode(gtf: Path):
    with gzip.open(gtf, "rt", encoding="utf-8") as fh:
        for line in fh:
            if line.startswith("#"):
                continue
            p = line.rstrip("\n").split("\t")
            if len(p) != 9 or p[2] != "transcript":
                continue
            a = dict(ATTR_RE.findall(p[8]))
            tid = strip_version(a.get("transcript_id", ""))
            gid = strip_version(a.get("gene_id", ""))
            start, end = int(p[3]), int(p[4])
            strand = p[6]
            tss = start if strand == "+" else end
            yield {
                "transcript_id": tid,
                "gene_id": gid,
                "gene_name": a.get("gene_name", ""),
                "transcript_type": a.get("transcript_type", a.get("transcript_biotype", "")),
                "chrom": p[0],
                "tss_1based": tss,
                "strand": strand,
                "candidate_promoter_id": f"GENCODE50:{tid}:{p[0]}:{tss}:{strand}",
            }


def load_bed_index(path: Path):
    arr = defaultdict(list)
    opener = gzip.open if str(path).endswith(".gz") else open
    with opener(path, "rt", encoding="utf-8") as fh:
        for line in fh:
            if not line.strip() or line.startswith("#"):
                continue
            p = line.rstrip("\n").split("\t")
            if len(p) < 3:
                continue
            arr[p[0]].append((int(p[1]), int(p[2]), p[3:]))
    for c in arr:
        arr[c].sort(key=lambda x: (x[0], x[1]))
    starts = {c: [x[0] for x in xs] for c, xs in arr.items()}
    return arr, starts


def point_hits(index, starts, chrom, pos1):
    x = pos1 - 1
    xs = index.get(chrom, [])
    ss = starts.get(chrom, [])
    i = bisect.bisect_right(ss, x)
    out = []
    j = i - 1
    while j >= 0 and xs[j][1] > x:
        if xs[j][0] <= x < xs[j][1]:
            out.append(xs[j])
        j -= 1
    return out


def load_dong_data7(path: Path):
    """Return chromosome-aware Dong records plus any-transcript presence.

    Data 7 contains 88 transcript IDs repeated across chrX/chrY. Therefore ENST
    alone is not a safe join key. GENCODE v50 transcript IDs are unique, but the
    evidence bridge must preserve Dong's chromosome-specific records.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    next(rows)
    header = [str(x).strip() for x in next(rows)]
    idx = {h: i for i, h in enumerate(header)}

    by_key = {}
    any_tid = set()
    multiplicity = defaultdict(int)
    for r in rows:
        tid = strip_version(str(r[idx["id"]]))
        chrom = str(r[idx["chr"]])
        any_tid.add(tid)
        multiplicity[tid] += 1
        by_key[(tid, chrom)] = {
            "dong_gene_id": strip_version(str(r[idx["gene_id"]])),
            "dong_chrom": chrom,
            "dong_tss_1based": int(r[idx["TSS"]]),
            "dong_strand": str(r[idx["strand"]]),
            "dong_internal_promoter": bool(r[idx["internalPromoter"]]),
            "dong_fivemost": bool(r[idx["fivemost"]]),
        }
    return by_key, any_tid, multiplicity


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def build(args):
    screen_idx, screen_starts = load_bed_index(args.screen_pls)

    fantom_idx = fantom_starts = None
    if args.fantom_bed:
        fantom_idx, fantom_starts = load_bed_index(args.fantom_bed)

    if args.dong_data7:
        dong_by_key, dong_any_tid, dong_multiplicity = load_dong_data7(args.dong_data7)
    else:
        dong_by_key, dong_any_tid, dong_multiplicity = {}, set(), {}

    args.output.parent.mkdir(parents=True, exist_ok=True)
    counts = defaultdict(int)
    genes = set()
    opener = gzip.open if str(args.output).endswith(".gz") else open

    fields = [
        "candidate_promoter_id",
        "transcript_id",
        "gene_id",
        "gene_name",
        "transcript_type",
        "chrom",
        "tss_1based",
        "strand",
        "screen_pls_overlap",
        "screen_pls_count",
        "dong_transcript_id_present_any",
        "dong_same_chrom_present",
        "dong_gene_match",
        "dong_coordinate_match",
        "dong_internal_promoter",
        "dong_fivemost",
        "fantom_cage_overlap",
        "fantom_cage_count",
    ]

    with opener(args.output, "wt", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=fields, delimiter="\t")
        w.writeheader()

        for g in load_gencode(args.gencode_gtf):
            counts["candidates"] += 1
            genes.add(g["gene_id"])

            sh = point_hits(screen_idx, screen_starts, g["chrom"], g["tss_1based"])
            dong_any = int(g["transcript_id"] in dong_any_tid)
            d = dong_by_key.get((g["transcript_id"], g["chrom"]))
            fhits = (
                point_hits(fantom_idx, fantom_starts, g["chrom"], g["tss_1based"])
                if fantom_idx is not None
                else []
            )

            row = {
                **g,
                "screen_pls_overlap": int(bool(sh)),
                "screen_pls_count": len(sh),
                "dong_transcript_id_present_any": dong_any,
                "dong_same_chrom_present": int(d is not None),
                "dong_gene_match": int(bool(d and d["dong_gene_id"] == g["gene_id"])),
                "dong_coordinate_match": int(
                    bool(
                        d
                        and d["dong_gene_id"] == g["gene_id"]
                        and d["dong_tss_1based"] == g["tss_1based"]
                        and d["dong_strand"] == g["strand"]
                    )
                ),
                "dong_internal_promoter": "" if d is None else int(d["dong_internal_promoter"]),
                "dong_fivemost": "" if d is None else int(d["dong_fivemost"]),
                "fantom_cage_overlap": "" if fantom_idx is None else int(bool(fhits)),
                "fantom_cage_count": "" if fantom_idx is None else len(fhits),
            }

            if sh:
                counts["screen_pls_overlap"] += 1
            if dong_any:
                counts["dong_transcript_id_present_any"] += 1
            if d:
                counts["dong_same_chrom_present"] += 1
            if row["dong_coordinate_match"]:
                counts["dong_coordinate_match"] += 1
            if fantom_idx is not None and fhits:
                counts["fantom_cage_overlap"] += 1

            w.writerow({k: row[k] for k in fields})

    summary = {
        "schema": "V64_PROMOTER_CANDIDATE_LEDGER_BUILD_RECEIPT_V1",
        "candidate_authority": "GENCODE_TRANSCRIPT_TSS",
        "candidate_join_key": "GENCODE transcript_id + chromosome + TSS + strand",
        "candidates": counts["candidates"],
        "genes": len(genes),
        "screen_pls_overlap": counts["screen_pls_overlap"],
        "dong_transcript_id_present_any": counts["dong_transcript_id_present_any"],
        "dong_same_chrom_present": counts["dong_same_chrom_present"],
        "dong_coordinate_match": counts["dong_coordinate_match"],
        "dong_duplicate_transcript_ids": sum(1 for v in dong_multiplicity.values() if v > 1),
        "fantom_coordinate_evidence_loaded": fantom_idx is not None,
        "fantom_cage_overlap": (
            counts["fantom_cage_overlap"] if fantom_idx is not None else None
        ),
        "inputs": {
            "gencode_gtf": {
                "path": str(args.gencode_gtf),
                "sha256": sha256_file(args.gencode_gtf),
            },
            "screen_pls": {
                "path": str(args.screen_pls),
                "sha256": sha256_file(args.screen_pls),
            },
            "dong_data7": (
                None
                if not args.dong_data7
                else {
                    "path": str(args.dong_data7),
                    "sha256": sha256_file(args.dong_data7),
                }
            ),
            "fantom_bed": (
                None
                if not args.fantom_bed
                else {
                    "path": str(args.fantom_bed),
                    "sha256": sha256_file(args.fantom_bed),
                }
            ),
        },
        "output": {
            "path": str(args.output),
            "sha256": sha256_file(args.output),
        },
        "denominator_rule": "No evidence layer is permitted to remove a GENCODE candidate.",
        "training_authorized": False,
    }
    args.receipt.write_text(
        json.dumps(summary, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    return summary


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--gencode-gtf", type=Path, required=True)
    ap.add_argument("--screen-pls", type=Path, required=True)
    ap.add_argument("--dong-data7", type=Path)
    ap.add_argument("--fantom-bed", type=Path)
    ap.add_argument("--output", type=Path, required=True)
    ap.add_argument("--receipt", type=Path, required=True)
    a = ap.parse_args(argv)
    print(json.dumps(build(a), indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
