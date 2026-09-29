#!/usr/bin/env python3
"""Apply the prospectively frozen V64 Nott C3 exact-identity successor rule.

This is an adjudicator, not a new liftover design. Authority:
  results/v64/V64_NOTT_C3_EXACT_IDENTITY_SUCCESSOR_RULE_V1.json

Frozen admissibility:
  * UCSC liftOver; minMatch=0.95
  * authoritative mapping without -multiple
  * separate -multiple pass only to detect/reject ambiguity
  * no nearest-neighbour rescue
  * each forward anchor: exactly one mapping, same chromosome, exact length
  * each retained pair: both anchors satisfy the forward rule
  * reverse lift with same semantics
  * each retained anchor returns exactly to original chr/start/end
  * final C3 pair: both anchors exact round-trip

All failures are attrition from the original 104,802 denominator.
"""
from __future__ import annotations
import argparse, hashlib, json, os, subprocess
from collections import Counter
from pathlib import Path

import openpyxl

EXPECTED = {
    "table_s5": "81c99689533d9da372cecdd469e7ff02cc985720105b83b3bd66c3ac8c93972e",
    "forward_chain": "5c0598e500ceb5a78c73086929e8ef993aec309bcafb595139b53d440b125a1d",
    "reverse_chain": "14a712e8e147d9fc8e9d87d51977b46f6f8ddb93efbe5d0843d86b6205f587b1",
    "liftover_v479": "80c77de53b8bbd5fec661242f24d4b2f0ac54446df6954d934d1a838927dd19c",
}
N_EXPECTED = 104802

def sha256(path):
    h=hashlib.sha256()
    with open(path,"rb") as f:
        for b in iter(lambda:f.read(1<<20), b""):
            h.update(b)
    return h.hexdigest()

def verify(path, expected, label):
    got=sha256(path)
    if got != expected:
        raise SystemExit(f"STOP_DIGEST_MISMATCH {label} got={got} expected={expected}")
    return got

def load_pairs(xlsx):
    wb=openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws=wb["Microglia interactome"]
    # The authenticated receipt fixed header_row_index=2 in pandas convention,
    # i.e. Excel row 3. Find chr1/start1/end1 robustly but require exact names.
    header=None
    rows=[]
    for row in ws.iter_rows(values_only=True):
        vals=[str(v) if v is not None else "" for v in row]
        if {"chr1","start1","end1","chr2","start2","end2"}.issubset(set(vals)):
            header=vals
            break
    if header is None:
        raise SystemExit("STOP_INTERACTOME_HEADER_NOT_FOUND")
    ix={k:header.index(k) for k in ["chr1","start1","end1","chr2","start2","end2"]}
    for row in ws.iter_rows(values_only=True):
        # iter_rows restarted? read_only worksheet iter can be restarted in openpyxl.
        vals=list(row)
        if len(vals) <= max(ix.values()):
            continue
        if vals[ix["chr1"]] == "chr1" and vals[ix["start1"]] == "start1":
            continue
        try:
            p=(str(vals[ix["chr1"]]), int(vals[ix["start1"]]), int(vals[ix["end1"]]),
               str(vals[ix["chr2"]]), int(vals[ix["start2"]]), int(vals[ix["end2"]]))
        except Exception:
            continue
        rows.append(p)
    if len(rows) != N_EXPECTED:
        raise SystemExit(f"STOP_DENOMINATOR_MISMATCH parsed={len(rows)} expected={N_EXPECTED}")
    if any((e1-s1)!=5000 or (e2-s2)!=5000 for c1,s1,e1,c2,s2,e2 in rows):
        raise SystemExit("STOP_NON_5000_SOURCE_ANCHOR")
    return rows

def run(cmd):
    cp=subprocess.run(cmd, text=True, capture_output=True)
    if cp.returncode != 0:
        raise SystemExit("STOP_LIFTOVER_FAILED\nCMD="+repr(cmd)+"\nOUT="+cp.stdout[-2000:]+"\nERR="+cp.stderr[-2000:])
    return {"cmd":cmd,"stdout":cp.stdout,"stderr":cp.stderr}

def read_bed(path):
    out={}
    dup=Counter()
    with open(path) as f:
        for line in f:
            if not line.strip() or line.startswith("#"): continue
            x=line.rstrip("\n").split("\t")
            if len(x)<4: continue
            name=x[3]
            dup[name]+=1
            if name not in out:
                out[name]=(x[0],int(x[1]),int(x[2]))
    return out,{k for k,v in dup.items() if v>1}

def multi_names(path):
    c=Counter()
    with open(path) as f:
        for line in f:
            if not line.strip() or line.startswith("#"): continue
            x=line.rstrip("\n").split("\t")
            if len(x)>=4: c[x[3]]+=1
    return {k for k,v in c.items() if v>1}

def unmapped_names(path):
    s=set(); reasons=Counter(); reason="UNKNOWN"
    with open(path) as f:
        for line in f:
            if line.startswith("#"):
                reason=line[1:].strip() or "UNKNOWN"
            elif line.strip():
                x=line.rstrip("\n").split("\t")
                if len(x)>=4:
                    s.add(x[3]); reasons[reason]+=1
    return s,dict(reasons)

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--table-s5",required=True)
    ap.add_argument("--forward-chain",required=True)
    ap.add_argument("--reverse-chain",required=True)
    ap.add_argument("--liftover",required=True)
    ap.add_argument("--out-dir",required=True)
    a=ap.parse_args()
    od=Path(a.out_dir)
    if od.exists() and any(od.iterdir()): raise SystemExit("STOP_OUTPUT_EXISTS")
    od.mkdir(parents=True,exist_ok=True)

    digests={
        "table_s5":verify(a.table_s5,EXPECTED["table_s5"],"table_s5"),
        "forward_chain":verify(a.forward_chain,EXPECTED["forward_chain"],"forward_chain"),
        "reverse_chain":verify(a.reverse_chain,EXPECTED["reverse_chain"],"reverse_chain"),
        "liftover_v479":verify(a.liftover,EXPECTED["liftover_v479"],"liftover_v479"),
    }
    pairs=load_pairs(a.table_s5)
    orig={}
    src=od/"anchors.hg19.bed"
    with open(src,"w") as f:
        for i,(c1,s1,e1,c2,s2,e2) in enumerate(pairs):
            for side,c,s,e in (("A",c1,s1,e1),("B",c2,s2,e2)):
                n=f"{side}{i}"
                orig[n]=(c,s,e)
                f.write(f"{c}\t{s}\t{e}\t{n}\n")

    fwd=od/"anchors.hg38.bed"; fwd_un=od/"anchors.hg38.unmapped"
    fwd_multi=od/"anchors.hg38.multi.bed"; fwd_multi_un=od/"anchors.hg38.multi.unmapped"
    c1=run([a.liftover,"-minMatch=0.95",str(src),a.forward_chain,str(fwd),str(fwd_un)])
    c2=run([a.liftover,"-minMatch=0.95","-multiple","-noSerial",str(src),a.forward_chain,str(fwd_multi),str(fwd_multi_un)])
    m,_=read_bed(fwd)
    amb=multi_names(fwd_multi)
    un,un_reasons=unmapped_names(fwd_un)

    status={}
    exact_forward={}
    for n,o in orig.items():
        if n in amb:
            status[n]="AMBIGUOUS"
        elif n not in m:
            status[n]="UNMAPPED"
        else:
            d=m[n]
            if d[0] != o[0]:
                status[n]="CHROMOSOME_CHANGED"
            elif (d[2]-d[1]) != (o[2]-o[1]):
                status[n]="LENGTH_CHANGED"
            else:
                status[n]="FORWARD_EXACT_LENGTH"
                exact_forward[n]=d

    fwd_exact_pairs=[]
    for i in range(N_EXPECTED):
        if f"A{i}" in exact_forward and f"B{i}" in exact_forward:
            fwd_exact_pairs.append(i)

    rt_in=od/"anchors.hg38.exact_forward.bed"
    with open(rt_in,"w") as f:
        for i in fwd_exact_pairs:
            for side in ("A","B"):
                n=f"{side}{i}"; c,s,e=exact_forward[n]
                f.write(f"{c}\t{s}\t{e}\t{n}\n")

    rt=od/"anchors.rt.hg19.bed"; rt_un=od/"anchors.rt.unmapped"
    rt_multi=od/"anchors.rt.multi.bed"; rt_multi_un=od/"anchors.rt.multi.unmapped"
    c3=run([a.liftover,"-minMatch=0.95",str(rt_in),a.reverse_chain,str(rt),str(rt_un)])
    c4=run([a.liftover,"-minMatch=0.95","-multiple","-noSerial",str(rt_in),a.reverse_chain,str(rt_multi),str(rt_multi_un)])
    rm,_=read_bed(rt)
    ramb=multi_names(rt_multi)
    runm,run_reasons=unmapped_names(rt_un)

    final_pairs=[]
    rt_anchor_counts=Counter()
    for i in fwd_exact_pairs:
        ok=True
        for side in ("A","B"):
            n=f"{side}{i}"
            if n in ramb:
                rt_anchor_counts["AMBIGUOUS"]+=1; ok=False
            elif n not in rm:
                rt_anchor_counts["UNMAPPED"]+=1; ok=False
            elif rm[n] != orig[n]:
                rt_anchor_counts["NONEXACT"]+=1; ok=False
            else:
                rt_anchor_counts["EXACT"]+=1
        if ok: final_pairs.append(i)

    fc=Counter(status.values())
    result={
        "schema":"V64_NOTT_C3_EXACT_IDENTITY_ADJUDICATION_V1",
        "date":"2026-09-29",
        "authority":{
            "rule":"results/v64/V64_NOTT_C3_EXACT_IDENTITY_SUCCESSOR_RULE_V1.json",
            "rule_commit":"38d789aceecfdc71beca8852fbb19ab1621fee55",
            "rule_status":"PROSPECTIVELY_FROZEN_BEFORE_C3_OUTCOME_ADJUDICATION",
            "claude_descriptive_evidence_commit":"5131bc6c44cae05f44ba8ef1f1605d429baf21ae"
        },
        "verified_input_sha256":digests,
        "source_pairs":N_EXPECTED,
        "forward_anchor_funnel":{
            "anchors_total":2*N_EXPECTED,
            "status_counts":dict(fc),
            "unmapped_reasons":un_reasons,
            "forward_exact_length_pairs":len(fwd_exact_pairs),
            "forward_exact_length_pair_fraction_original":len(fwd_exact_pairs)/N_EXPECTED,
        },
        "roundtrip_on_forward_exact_pairs":{
            "anchors_attempted":2*len(fwd_exact_pairs),
            "anchor_status_counts":dict(rt_anchor_counts),
            "unmapped_reasons":run_reasons,
            "exact_both_anchor_pairs":len(final_pairs),
            "exact_both_anchor_fraction_forward_exact":len(final_pairs)/max(1,len(fwd_exact_pairs)),
            "exact_both_anchor_fraction_original":len(final_pairs)/N_EXPECTED,
        },
        "C3":{
            "PASS":bool(final_pairs),
            "qualification":"PASS" if final_pairs else "FAIL",
            "qualified_pair_count":len(final_pairs),
            "qualified_pair_fraction_original":len(final_pairs)/N_EXPECTED,
            "interpretation":"PASS means every retained pair satisfies the pre-frozen exact forward-length and exact round-trip identity rules. It does not establish biological substrate fit or support representativeness."
        },
        "commands":[c1["cmd"],c2["cmd"],c3["cmd"],c4["cmd"]],
        "governance":{"training":"OFF","td60":"BLOCKED","biological_outcomes_opened":False}
    }
    with open(od/"V64_NOTT_C3_EXACT_IDENTITY_ADJUDICATION_V1.json","w") as f:
        json.dump(result,f,indent=2)
    with open(od/"V64_NOTT_C3_QUALIFIED_PAIR_INDICES_V1.txt","w") as f:
        for i in final_pairs: f.write(str(i)+"\n")
    print(json.dumps(result,indent=2))
if __name__=="__main__": main()
