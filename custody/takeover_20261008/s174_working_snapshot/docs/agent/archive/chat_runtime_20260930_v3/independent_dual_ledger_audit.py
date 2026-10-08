import gzip,re,bisect,csv,hashlib,json
from pathlib import Path
from collections import defaultdict,Counter
from openpyxl import load_workbook
BASE=Path('/mnt/data/openres_art/open_resource_cache_v1')
OUT=Path('/mnt/data/dual_ledger_independent'); OUT.mkdir(exist_ok=True)
GTF=BASE/'GENCODE_V50_COMPREHENSIVE_CHR_GTF__01__gencode.v50.annotation.gtf.gz'
SCREEN=BASE/'SCREEN_REGISTRY_V4_GRCH38_PLS__01__GRCh38-cCREs.PLS.bed'
FANTOM=BASE/'FANTOM5_HG38_CAGE_PEAK_COORDINATES__01__hg38_fair+new_CAGE_peaks_phase1and2.bed.gz'
DONG=BASE/'DONG_ROUSSOS_2024_AUTOMATED_SUPPLEMENTS_7_9__01__41467_2024_54448_MOESM9_ESM.xlsx'
ATTR_RE=re.compile(r'(\S+) "([^"]+)"')
def stripv(x): return x.split('.',1)[0] if x else x
def sha(p):
 h=hashlib.sha256();
 with open(p,'rb') as f:
  for ch in iter(lambda:f.read(1<<20),b''): h.update(ch)
 return h.hexdigest()
def load_bed(path):
 arr=defaultdict(list); op=gzip.open if str(path).endswith('.gz') else open
 with op(path,'rt') as f:
  for line in f:
   if not line.strip() or line.startswith('#'): continue
   p=line.rstrip('\n').split('\t'); arr[p[0]].append((int(p[1]),int(p[2]),p[3:]))
 starts={}; pmax={}
 for c,xs in arr.items():
  xs.sort(key=lambda x:(x[0],x[1])); starts[c]=[x[0] for x in xs]; cur=-1; pm=[]
  for _,e,_ in xs: cur=max(cur,e); pm.append(cur)
  pmax[c]=pm
 return arr,starts,pmax
def hits(idx,starts,pmax,chrom,pos1):
 x=pos1-1; xs=idx.get(chrom,[]); ss=starts.get(chrom,[]); pm=pmax.get(chrom,[]); j=bisect.bisect_right(ss,x)-1; out=[]
 while j>=0 and pm[j]>x:
  if xs[j][0]<=x<xs[j][1]: out.append(xs[j])
  j-=1
 return out
# Dong
wb=load_workbook(DONG,read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]; it=ws.iter_rows(values_only=True); next(it); hdr=[str(x).strip() for x in next(it)]; ix={h:i for i,h in enumerate(hdr)}
dong={}; anytid=set(); mult=Counter()
for r in it:
 tid=stripv(str(r[ix['id']])); chrom=str(r[ix['chr']]); anytid.add(tid); mult[tid]+=1
 dong[(tid,chrom)]={'gene_id':stripv(str(r[ix['gene_id']])), 'tss':int(r[ix['TSS']]), 'strand':str(r[ix['strand']]), 'internal':bool(r[ix['internalPromoter']]), 'five':bool(r[ix['fivemost']])}
si,ss,sm=load_bed(SCREEN); fi,fs,fm=load_bed(FANTOM)
TF=OUT/'PROMOTER_TRANSCRIPT_LEDGER.tsv.gz'; MF=OUT/'TRANSCRIPT_TO_TSS_MEMBERSHIP.tsv.gz'; XF=OUT/'EXACT_TSS_LEDGER.tsv.gz'
tfields=['candidate_promoter_id','exact_tss_id','transcript_id','gene_id','gene_name','transcript_type','chrom','tss_1based','strand','screen_pls_overlap','screen_pls_count','dong_transcript_id_present_any','dong_same_chrom_present','dong_gene_match','dong_coordinate_match','dong_internal_promoter','dong_fivemost','fantom_cage_overlap','fantom_cage_count','fantom_representative_tss_exact']
mfields=['candidate_promoter_id','transcript_id','gene_id','exact_tss_id','chrom','tss_1based','strand']
tss={}; genes=set(); tx=0; dong_counts=Counter()
with gzip.open(TF,'wt',newline='') as tfh,gzip.open(MF,'wt',newline='') as mfh:
 tw=csv.DictWriter(tfh,fieldnames=tfields,delimiter='\t'); mw=csv.DictWriter(mfh,fieldnames=mfields,delimiter='\t'); tw.writeheader(); mw.writeheader()
 with gzip.open(GTF,'rt') as f:
  for line in f:
   if line.startswith('#'): continue
   p=line.rstrip('\n').split('\t')
   if len(p)!=9 or p[2]!='transcript': continue
   a=dict(ATTR_RE.findall(p[8])); tid=stripv(a.get('transcript_id','')); gid=stripv(a.get('gene_id','')); strand=p[6]; start,end=int(p[3]),int(p[4]); pos=start if strand=='+' else end; chrom=p[0]
   tx+=1; genes.add(gid); xid=f'GENCODE50:TSS:{gid}:{chrom}:{pos}:{strand}'; cid=f'GENCODE50:{tid}:{chrom}:{pos}:{strand}'
   sh=hits(si,ss,sm,chrom,pos); fh=[h for h in hits(fi,fs,fm,chrom,pos) if len(h[2])>=3 and h[2][2]==strand]; d=dong.get((tid,chrom)); exact=int(any(len(h[2])>=4 and int(h[2][3])==pos-1 for h in fh))
   row={'candidate_promoter_id':cid,'exact_tss_id':xid,'transcript_id':tid,'gene_id':gid,'gene_name':a.get('gene_name',''),'transcript_type':a.get('transcript_type',a.get('transcript_biotype','')),'chrom':chrom,'tss_1based':pos,'strand':strand,'screen_pls_overlap':int(bool(sh)),'screen_pls_count':len(sh),'dong_transcript_id_present_any':int(tid in anytid),'dong_same_chrom_present':int(d is not None),'dong_gene_match':int(bool(d and d['gene_id']==gid)),'dong_coordinate_match':int(bool(d and d['gene_id']==gid and d['tss']==pos and d['strand']==strand)),'dong_internal_promoter':'' if d is None else int(d['internal']),'dong_fivemost':'' if d is None else int(d['five']),'fantom_cage_overlap':int(bool(fh)),'fantom_cage_count':len(fh),'fantom_representative_tss_exact':exact}
   tw.writerow(row); mw.writerow({k:row[k] for k in mfields})
   for k in ['dong_transcript_id_present_any','dong_same_chrom_present','dong_gene_match','dong_coordinate_match']: dong_counts[k]+=row[k]
   if xid not in tss:
    tss[xid]={'exact_tss_id':xid,'gene_id':gid,'gene_name':a.get('gene_name',''),'chrom':chrom,'tss_1based':pos,'strand':strand,'transcript_count':0,'screen_pls_overlap':int(bool(sh)),'screen_pls_count':len(sh),'fantom_cage_overlap':int(bool(fh)),'fantom_cage_count':len(fh),'fantom_representative_tss_exact':exact,'dong_transcripts_present_any':0,'dong_same_chrom_transcripts':0,'dong_coordinate_match_transcripts':0}
   z=tss[xid]; z['transcript_count']+=1; z['dong_transcripts_present_any']+=row['dong_transcript_id_present_any']; z['dong_same_chrom_transcripts']+=row['dong_same_chrom_present']; z['dong_coordinate_match_transcripts']+=row['dong_coordinate_match']
xfields=['exact_tss_id','gene_id','gene_name','chrom','tss_1based','strand','transcript_count','screen_pls_overlap','screen_pls_count','fantom_cage_overlap','fantom_cage_count','fantom_representative_tss_exact','dong_transcripts_present_any','dong_same_chrom_transcripts','dong_coordinate_match_transcripts']
with gzip.open(XF,'wt',newline='') as fh:
 w=csv.DictWriter(fh,fieldnames=xfields,delimiter='\t'); w.writeheader()
 for xid in sorted(tss): w.writerow(tss[xid])
sc=sum(x['screen_pls_overlap'] for x in tss.values()); fc=sum(x['fantom_cage_overlap'] for x in tss.values()); ex=sum(x['fantom_representative_tss_exact'] for x in tss.values()); joint=Counter()
for x in tss.values(): joint[('S' if x['screen_pls_overlap'] else 's')+('F' if x['fantom_cage_overlap'] else 'f')]+=1
rec={'transcript_records':tx,'genes':len(genes),'exact_tss_loci':len(tss),'screen_tss':sc,'fantom_tss':fc,'fantom_exact':ex,'joint':dict(joint),'dong_transcript_counts':dict(dong_counts),'dong_duplicate_ids':sum(v>1 for v in mult.values()),'outputs':{p.name:{'bytes':p.stat().st_size,'sha256':sha(p)} for p in [TF,XF,MF]}}
(OUT/'INDEPENDENT_RECEIPT.json').write_text(json.dumps(rec,indent=2,sort_keys=True)+'\n')
print(json.dumps(rec,indent=2,sort_keys=True))
