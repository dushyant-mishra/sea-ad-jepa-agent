#!/usr/bin/env python3
"""P1S V2 — contract-compliant DESCRIPTIVE SENSITIVITY rerun. Primary untouched.

Narrow scope, per the audit at 342d5796:
  PRIMARY_GATE_PASS_VERIFIED__FULL_CONTRACT_CLOSURE_PENDING_SENSITIVITY_RERUN

Two compliance defects in my first runner, both accepted:

  1. LOCAL DENSITY. V2 freezes: "For each observed/null distal interval AND EACH
     ATAC TRACK, count peak midpoints within +/-50 kb of the interval midpoint …
     Report E_c and D contrasts by POOLED quartiles of local peak density."
     My runner binned on MICROGLIA OBSERVED-distal density alone -- one track,
     one interval class. That is not the pooled quantity the contract names.

  2. OMITTED MANDATORY REPORTS. V2's mandatory_reporting requires "results by
     predeclared quartiles of post-C3 log distance and promoter degree as
     descriptive sensitivity only". My first run produced neither.

WHAT THIS SCRIPT DOES NOT TOUCH: population, null construction, estimand,
bootstrap, thresholds, or the PASS decision. It recomputes the identical analysed
population and ASSERTS the primary quantities reproduce exactly before emitting
any sensitivity table; if they do not, it stops rather than reporting.

POOLED DENSITY, interpretation stated rather than assumed. The contract defines
density per (interval, track) and asks for "pooled quartiles". The literal pooled
quantity per analysed pair is therefore the sum over all three tracks AND over
both the observed and the null distal interval. Pooling across observed and null
matters: it makes the stratum assignment independent of which interval happens to
be the observed one, so a pair cannot be sorted into a stratum by the very
quantity under test. The superseded microglia-only stratification is reported
alongside, clearly labelled, so the earlier Q1 number stays traceable.

TRAINING=OFF. TD60=BLOCKED.
"""
from __future__ import annotations

import argparse
import gzip
import hashlib
import json
import os
import sys
from bisect import bisect_right
from collections import Counter, defaultdict

import numpy as np
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import importlib.util

HERE = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location(
    "p1s", os.path.join(HERE, "nott_p1s_v2_substrate_fit_v1.py"))
P = importlib.util.module_from_spec(spec)
spec.loader.exec_module(P)

PRIMARY = "results/v64/V64_NOTT_P1S_V2_RESULT_V1.json"


def quart_labels(v, k=4):
    q = np.quantile(v, [0.25, 0.5, 0.75])
    return np.digitize(v, q), q


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out-dir", required=True)
    a = ap.parse_args()
    if os.path.exists(a.out_dir) and os.listdir(a.out_dir):
        raise SystemExit("STOP_OUTPUT_EXISTS")
    os.makedirs(a.out_dir, exist_ok=True)
    if P.sha(P.S5) != P.S5_SHA:
        raise SystemExit("STOP_S5_DIGEST_MISMATCH")

    # ---- rebuild the identical analysed population (same code paths as primary)
    wb = openpyxl.load_workbook(P.S5, read_only=True, data_only=True)
    ws = wb["H3K4me3_around_TSS_annotated_pe"]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows)
              if r and any(str(x).strip() == "PeakID" for x in r if x))
    hdr = [str(x).strip() if x is not None else "" for x in rows[hi]]
    col = {k: i for i, k in enumerate(hdr)}
    prom_by_chrom = defaultdict(list)
    for r in rows[hi + 1:]:
        if not r or r[col["Chr"]] is None:
            continue
        flag = r[col["PU1_active_promoter"]]
        if not (flag is True or str(flag).strip().upper() in ("TRUE", "1", "YES")):
            continue
        ens = r[col["Nearest Ensembl"]]
        prom_by_chrom[str(r[col["Chr"]])].append(
            (int(r[col["Start"]]), int(r[col["End"]]),
             str(ens).strip() if ens is not None else ""))
    for c in prom_by_chrom:
        prom_by_chrom[c].sort()

    ws = wb["Microglia interactome"]
    irows = list(ws.iter_rows(values_only=True))
    ih = next(i for i, r in enumerate(irows) if r and r[0] == "chr1")
    inter = [(str(r[0]), int(r[1]), int(r[2]), str(r[3]), int(r[4]), int(r[5]))
             for r in irows[ih + 1:] if r and r[0] is not None]

    def prom_hits(c, s, e):
        out = []
        for ps, pe, ens in prom_by_chrom.get(c, []):
            if ps >= e:
                break
            if pe > s:
                out.append(ens)
        return out

    orient = {}
    for i, (c1, s1, e1, c2, s2, e2) in enumerate(inter):
        h1, h2 = prom_hits(c1, s1, e1), prom_hits(c2, s2, e2)
        p1, p2 = len(h1) > 0, len(h2) > 0
        if p1 and p2 or (not p1 and not p2):
            orient[i] = (None, None)
        else:
            side = "A" if p1 else "B"
            genes = {g for g in (h1 if p1 else h2) if g}
            k = ("PROMOTER_DISTAL_UNIQUE_ENSEMBL" if len(genes) == 1
                 else "GENE_AMBIGUOUS_PROMOTER_DISTAL" if len(genes) > 1
                 else "GENE_UNANNOTATED_PROMOTER_DISTAL")
            orient[i] = (k, side)

    cur = defaultdict(dict)
    with gzip.open(P.DISPO, "rt") as fh:
        h = fh.readline().rstrip("\n").split("\t")
        ci = {k: j for j, k in enumerate(h)}
        for line in fh:
            f = line.rstrip("\n").split("\t")
            if f[ci["pair_C3_retained"]] != "True":
                continue
            cur[int(f[ci["pair_index"]])][f[ci["side"]]] = (
                f[ci["hg38_chrom"]], int(f[ci["hg38_start"]]), int(f[ci["hg38_end"]]))
    c3 = {i: d for i, d in cur.items() if "A" in d and "B" in d}

    OK = {"PROMOTER_DISTAL_UNIQUE_ENSEMBL", "GENE_AMBIGUOUS_PROMOTER_DISTAL",
          "GENE_UNANNOTATED_PROMOTER_DISTAL"}
    pop = [i for i in c3 if orient[i][0] in OK]

    tracks = {}
    for name, path in P.ATAC.items():
        iv = P.Intervals()
        with gzip.open(path, "rt") as fh:
            for line in fh:
                if line.startswith("#") or not line.strip():
                    continue
                f = line.rstrip("\n").split("\t")
                iv.add(f[0], int(f[1]), int(f[2]))
        iv.build()
        tracks[name] = iv
    sizes = P.hg38_chrom_sizes(P.CHAIN)
    mids = {t: {c: np.array(sorted((s + e) // 2 for s, e in iv.d[c]))
                for c in iv.d} for t, iv in tracks.items()}

    rec = []
    for i in sorted(pop):
        side = orient[i][1]
        pc, ps, pe = c3[i][side]
        dc, ds, de = c3[i]["B" if side == "A" else "A"]
        if pc != dc:
            continue
        pm, dm, w = (ps + pe) // 2, (ds + de) // 2, de - ds
        d = dm - pm
        nm = pm - d
        ns, ne = nm - w // 2, nm - w // 2 + w
        if ns < 0 or ne > sizes.get(dc, 0):
            continue
        r = {"i": i, "prom": (pc, ps, pe), "chrom": dc,
             "obs_mid": dm, "null_mid": (ns + ne) // 2, "abs_d": abs(d)}
        for t, iv in tracks.items():
            r["O_" + t] = int(iv.overlaps(dc, ds, de))
            r["N_" + t] = int(iv.overlaps(dc, ns, ne))
        rec.append(r)

    n = len(rec)
    E = {t: np.array([x["O_" + t] - x["N_" + t] for x in rec], float) for t in tracks}
    D_MN = E["microglia"] - E["neuron"]
    D_MO = E["microglia"] - E["oligo"]

    # ---- FAIL-CLOSED: the primary must reproduce exactly before any table is emitted
    prim = json.load(open(PRIMARY))
    checks = {
        "analysed_pairs": bool(n == prim["population_funnel"]["analysed_pairs"]),
        "E_microglia_mean": bool(abs(E["microglia"].mean() - prim["primary"]["E_microglia"]["mean"]) < 1e-12),
        "D_MN_mean": bool(abs(D_MN.mean() - prim["primary"]["D_MN"]["mean"]) < 1e-12),
        "D_MO_mean": bool(abs(D_MO.mean() - prim["primary"]["D_MO"]["mean"]) < 1e-12),
    }
    if not all(checks.values()):
        raise SystemExit(f"STOP_PRIMARY_NOT_REPRODUCED {checks}")

    # ---- density per (interval, track), exactly as the contract defines it
    def dens(chrom, mid, t):
        arr = mids[t].get(chrom, np.array([]))
        return int(np.searchsorted(arr, mid + 50000) - np.searchsorted(arr, mid - 50000))

    per = {f"{t}_{w}": np.array([dens(x["chrom"], x[f"{w}_mid"], t) for x in rec], float)
           for t in tracks for w in ("obs", "null")}
    pooled = sum(per.values())                       # all 3 tracks x {observed, null}
    superseded_mic_obs = per["microglia_obs"]        # what the first run used

    def table(strata, q):
        out = {}
        for s in range(4):
            m = strata == s
            if m.sum() < 2:
                continue
            out[f"Q{s+1}"] = {"n": int(m.sum()),
                              "E_microglia": float(E["microglia"][m].mean()),
                              "E_neuron": float(E["neuron"][m].mean()),
                              "E_oligo": float(E["oligo"][m].mean()),
                              "D_MN": float(D_MN[m].mean()),
                              "D_MO": float(D_MO[m].mean())}
        out["_quartile_breaks"] = [float(x) for x in q]
        return out

    s_pool, q_pool = quart_labels(pooled)
    s_sup, q_sup = quart_labels(superseded_mic_obs)
    logd = np.log10(np.maximum(np.array([x["abs_d"] for x in rec], float), 1.0))
    s_dist, q_dist = quart_labels(logd)
    deg_of = Counter(x["prom"] for x in rec)
    degree = np.array([deg_of[x["prom"]] for x in rec], float)
    s_deg, q_deg = quart_labels(degree)

    out = {
        "schema": "V64_NOTT_P1S_V2_SENSITIVITY_RERUN_V1",
        "date": "2026-09-29",
        "scope": "DESCRIPTIVE SENSITIVITY ONLY. Population, null, estimand, bootstrap, "
                 "thresholds and the PASS decision are untouched and were verified to "
                 "reproduce exactly before any table below was emitted.",
        "responds_to_audit": "342d5796491f31fb1f47340a7bbb9e185bbc8db7 — "
                             "PRIMARY_GATE_PASS_VERIFIED__FULL_CONTRACT_CLOSURE_PENDING_SENSITIVITY_RERUN",
        "defects_corrected": [
            "local density was binned on MICROGLIA OBSERVED-distal density alone; the "
            "contract freezes POOLED quartiles over each observed/null interval and each "
            "ATAC track",
            "the mandatory post-C3 log-distance and promoter-degree quartile sensitivities "
            "were omitted entirely"],
        "primary_reproduction_check": checks,
        "primary_unchanged": {
            "E_microglia_mean": float(E["microglia"].mean()),
            "D_MN_mean": float(D_MN.mean()), "D_MO_mean": float(D_MO.mean()),
            "analysed_pairs": n,
            "R1_R2_R3": "PASS (unchanged; not recomputed here)"},
        "pooled_density_interpretation": (
            "The contract defines density per (interval, track) and asks for POOLED "
            "quartiles. The pooled quantity per pair is therefore the sum over all three "
            "ATAC tracks AND over both the observed and the null distal interval. Pooling "
            "across observed and null is deliberate: it makes stratum assignment "
            "independent of which interval is the observed one, so a pair cannot be "
            "sorted into a stratum by the quantity under test."),
        "SENSITIVITY_pooled_local_density": table(s_pool, q_pool),
        "SENSITIVITY_post_C3_log_distance": table(s_dist, q_dist),
        "SENSITIVITY_promoter_degree": table(s_deg, q_deg),
        "SUPERSEDED_microglia_observed_only_density": {
            "status": "NON-COMPLIANT with V2; retained only so the earlier reported Q1 "
                      "figure stays traceable. Not authoritative.",
            **table(s_sup, q_sup)},
        "per_interval_track_density_means": {k: float(v.mean()) for k, v in per.items()},
        "governance": {"training": "OFF", "td60": "BLOCKED",
                       "population_changed": False, "null_changed": False,
                       "estimand_changed": False, "thresholds_changed": False,
                       "pass_decision_changed": False},
    }
    with open(os.path.join(a.out_dir, "V64_NOTT_P1S_V2_SENSITIVITY_RERUN_V1.json"), "w") as fh:
        json.dump(out, fh, indent=2)

    print("primary reproduction:", checks)
    for key in ("SENSITIVITY_pooled_local_density", "SENSITIVITY_post_C3_log_distance",
                "SENSITIVITY_promoter_degree", "SUPERSEDED_microglia_observed_only_density"):
        print(f"\n{key}")
        for q, v in out[key].items():
            if q.startswith("_") or q == "status":
                continue
            print(f"   {q}: n={v['n']:>6}  E_mic {v['E_microglia']:+.4f}  "
                  f"D_MN {v['D_MN']:+.4f}  D_MO {v['D_MO']:+.4f}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
